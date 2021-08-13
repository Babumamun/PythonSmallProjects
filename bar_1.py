import openpyxl
from openpyxl import Workbook
import openpyxl as openpyxl
from openpyxl.chart import BarChart

wb = openpyxl.load_workbook('/Users/mac/Desktop/stu_scores _Grade 2.xlsx')
sheet = wb['stu_scores_01']

data = openpyxl.chart.Reference(sheet, min_col=3, min_row=34, max_row=34,max_col=7)
cat = openpyxl.chart.Reference(sheet, min_col=3, min_row=1, max_col=7,max_row=1)

seriesObj = openpyxl.chart.Series(data, title="bar chart of each subject", title_from_data=True)
charObj = openpyxl.chart.BarChart()
charObj.title = "My Bar Chart"
charObj.x_axis.title = 'bar chart of each subject'
charObj.append(seriesObj)
charObj.set_categories(cat)
sheet.add_chart(charObj,"I2")
#new one
data = openpyxl.chart.Reference(sheet, min_col=3, min_row=35, max_row=35,max_col=7)
cat = openpyxl.chart.Reference(sheet, min_col=3, min_row=1, max_col=7,max_row=1)

seriesObj = openpyxl.chart.Series(data, title="bar chart of boys  each subject", title_from_data=True)
charObj = openpyxl.chart.BarChart()
charObj.title = "My Bar Chart"
charObj.x_axis.title = 'Boys each subject'
charObj.append(seriesObj)
charObj.set_categories(cat)
sheet.add_chart(charObj,"I18")

data = openpyxl.chart.Reference(sheet, min_col=3, min_row=36, max_row=36,max_col=7)
cat = openpyxl.chart.Reference(sheet, min_col=3, min_row=1, max_col=7,max_row=1)

seriesObj = openpyxl.chart.Series(data, title="bar chart of girls each subject", title_from_data=True)
charObj = openpyxl.chart.BarChart()
charObj.title = "My Bar Chart"
charObj.x_axis.title = 'girls each subject'
charObj.append(seriesObj)
charObj.set_categories(cat)
sheet.add_chart(charObj,"Q2")

wb.save('stu_scores _Grade.xlsx')
wb.close()
